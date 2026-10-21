#!/usr/bin/env python3
"""
Trial 3 — NIFTY 500 end-of-day qualifying-stock scanner.

Conditions:
Monthly:
 M1: RSI(5) on last CLOSED monthly candle > SMA(14) of monthly RSI(5).
 M2: ADX(14) on last CLOSED monthly candle > ADX(14) on second-last CLOSED monthly candle.
 M3: 0 < (last-closed monthly ADX(14) - second-last-closed monthly ADX(14)) < 1.
Weekly:
 W1: RSI(5) on last CLOSED weekly candle > SMA(14) of weekly RSI(5).
 W2: MACD(12,26,9) line > signal on the last CLOSED weekly candle.
Daily:
 D1: daily low <= daily EMA(20) AND daily close > daily EMA(20).
 D2: daily MACD(12,26,9) line > signal on the daily candle being evaluated.

This is a qualifying-stock list only. It does not simulate entries, stops, targets, or exits.
"""

from __future__ import annotations

import io
import os
import sys
import time
import warnings
from datetime import datetime
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import requests
import yfinance as yf
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore", category=FutureWarning)

IST = ZoneInfo("Asia/Kolkata")
OUTPUT_XLSX = "trial3_qualifying_stocks.xlsx"
OUTPUT_CSV = "trial3_qualifying_stocks.csv"
CONSTITUENTS_URL = "https://www.niftyindices.com/IndexConstituent/ind_nifty500list.csv"
BACKUP_CONSTITUENTS_URL = "https://archives.nseindia.com/content/indices/ind_nifty500list.csv"
START_DATE = "2015-01-01"  # indicator warm-up for monthly/weekly calculations
DOWNLOAD_PERIOD = "max"


def get_constituents() -> pd.DataFrame:
    headers = {
        "User-Agent": "Mozilla/5.0",
        "Accept": "text/csv,*/*",
        "Referer": "https://www.niftyindices.com/",
    }
    last_error = None
    for url in (CONSTITUENTS_URL, BACKUP_CONSTITUENTS_URL):
        try:
            r = requests.get(url, headers=headers, timeout=25)
            r.raise_for_status()
            df = pd.read_csv(io.BytesIO(r.content))
            df.columns = [str(c).strip() for c in df.columns]
            symbol_col = next((c for c in df.columns if c.lower() in ("symbol", "ticker")), None)
            name_col = next((c for c in df.columns if c.lower() in ("company name", "companyname", "name")), None)
            if not symbol_col:
                raise ValueError(f"Could not find Symbol column in constituent file: {df.columns.tolist()}")
            out = pd.DataFrame()
            out["Symbol"] = df[symbol_col].astype(str).str.strip()
            out["Company"] = df[name_col].astype(str).str.strip() if name_col else out["Symbol"]
            out = out[out["Symbol"].ne("") & out["Symbol"].ne("NAN")].drop_duplicates("Symbol")
            if len(out) < 400:
                raise ValueError(f"Constituent file returned only {len(out)} symbols; refusing to scan incomplete universe.")
            out["Ticker"] = out["Symbol"] + ".NS"
            return out.reset_index(drop=True)
        except Exception as e:
            last_error = e
    raise RuntimeError(f"Could not download NIFTY 500 constituents. Last error: {last_error}")


def rsi(series: pd.Series, period: int = 5) -> pd.Series:
    delta = series.diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)
    # Wilder's smoothing, as commonly used for RSI.
    avg_gain = gain.ewm(alpha=1 / period, min_periods=period, adjust=False).mean()
    avg_loss = loss.ewm(alpha=1 / period, min_periods=period, adjust=False).mean()
    rs = avg_gain / avg_loss.replace(0, np.nan)
    result = 100 - (100 / (1 + rs))
    result = result.where(avg_loss.ne(0), 100)
    return result


def macd(close: pd.Series, fast: int = 12, slow: int = 26, signal: int = 9):
    line = close.ewm(span=fast, adjust=False, min_periods=fast).mean() - close.ewm(
        span=slow, adjust=False, min_periods=slow
    ).mean()
    sig = line.ewm(span=signal, adjust=False, min_periods=signal).mean()
    return line, sig


def adx(high: pd.Series, low: pd.Series, close: pd.Series, period: int = 14) -> pd.Series:
    prev_close = close.shift(1)
    up_move = high.diff()
    down_move = -low.diff()
    plus_dm = pd.Series(np.where((up_move > down_move) & (up_move > 0), up_move, 0.0), index=high.index)
    minus_dm = pd.Series(np.where((down_move > up_move) & (down_move > 0), down_move, 0.0), index=high.index)
    tr = pd.concat(
        [(high - low), (high - prev_close).abs(), (low - prev_close).abs()], axis=1
    ).max(axis=1)
    atr = tr.ewm(alpha=1 / period, min_periods=period, adjust=False).mean()
    plus_di = 100 * plus_dm.ewm(alpha=1 / period, min_periods=period, adjust=False).mean() / atr
    minus_di = 100 * minus_dm.ewm(alpha=1 / period, min_periods=period, adjust=False).mean() / atr
    denom = (plus_di + minus_di).replace(0, np.nan)
    dx = 100 * (plus_di - minus_di).abs() / denom
    return dx.ewm(alpha=1 / period, min_periods=period, adjust=False).mean()


def ohlc_resample(daily: pd.DataFrame, rule: str) -> pd.DataFrame:
    return daily.resample(rule, label="right", closed="right").agg(
        {"Open": "first", "High": "max", "Low": "min", "Close": "last", "Volume": "sum"}
    ).dropna(subset=["Open", "High", "Low", "Close"])


def completed_weekly_bars(daily: pd.DataFrame) -> pd.DataFrame:
    weekly = ohlc_resample(daily, "W-FRI")
    if weekly.empty:
        return weekly
    latest_date = daily.index.max().date()
    # A weekly candle is considered closed only after Friday's daily candle is present.
    # If the last trading session is earlier in the week (holiday or mid-week scan),
    # conservatively use the prior completed weekly candle.
    if latest_date.weekday() != 4 and weekly.index.max().date() >= latest_date:
        weekly = weekly.iloc[:-1]
    elif latest_date.weekday() != 4:
        # Current partial week has a Friday label later than the latest daily bar.
        weekly = weekly.loc[weekly.index.date < weekly.index.max().date()]
    return weekly


def completed_monthly_bars(daily: pd.DataFrame) -> pd.DataFrame:
    monthly = ohlc_resample(daily, "ME")
    if monthly.empty:
        return monthly
    latest = daily.index.max()
    # Always exclude the current calendar month. This is deliberately conservative:
    # it prevents a partial month from leaking into signals on any date in that month.
    current_month_start = pd.Timestamp(latest.year, latest.month, 1)
    monthly = monthly.loc[monthly.index < current_month_start]
    return monthly


def normalize_history(raw: pd.DataFrame, ticker: str) -> pd.DataFrame:
    if raw is None or raw.empty:
        return pd.DataFrame()
    # yfinance multi-ticker/single-ticker layouts vary by version.
    if isinstance(raw.columns, pd.MultiIndex):
        if ticker in raw.columns.get_level_values(-1):
            raw = raw.xs(ticker, axis=1, level=-1, drop_level=True)
        else:
            raw.columns = raw.columns.get_level_values(0)
    needed = ["Open", "High", "Low", "Close", "Volume"]
    if not all(c in raw.columns for c in needed):
        return pd.DataFrame()
    df = raw[needed].copy().dropna(subset=["Open", "High", "Low", "Close"])
    idx = pd.to_datetime(df.index)
    if getattr(idx, "tz", None) is not None:
        idx = idx.tz_localize(None)
    df.index = idx.normalize()
    df = df[~df.index.duplicated(keep="last")].sort_index()
    return df


def latest(series: pd.Series):
    return float(series.iloc[-1]) if len(series) and pd.notna(series.iloc[-1]) else np.nan


def evaluate_stock(symbol: str, company: str, ticker: str, raw: pd.DataFrame, scan_date: str):
    daily = normalize_history(raw, ticker)
    if daily.empty or len(daily) < 100:
        return None, "Insufficient daily history"

    # Remove data after the intended scan date, if the provider returned any.
    scan_day = pd.Timestamp(scan_date).normalize()
    daily = daily.loc[daily.index <= scan_day]
    if daily.empty:
        return None, "No data on or before scan date"

    # Daily indicators: use the most recent available daily bar (scan runs after close).
    d_close = daily["Close"]
    d_ema20 = d_close.ewm(span=20, adjust=False, min_periods=20).mean()
    d_macd, d_signal = macd(d_close)
    daily_row = {
        "Daily Date": daily.index[-1].date().isoformat(),
        "Close": latest(d_close),
        "EMA20": latest(d_ema20),
        "Daily Low": latest(daily["Low"]),
        "Daily MACD": latest(d_macd),
        "Daily MACD Signal": latest(d_signal),
    }
    d1 = pd.notna(daily_row["EMA20"]) and daily_row["Daily Low"] <= daily_row["EMA20"] and daily_row["Close"] > daily_row["EMA20"]
    d2 = pd.notna(daily_row["Daily MACD"]) and pd.notna(daily_row["Daily MACD Signal"]) and daily_row["Daily MACD"] > daily_row["Daily MACD Signal"]

    # Weekly indicators: use completed weekly candles only.
    weekly = completed_weekly_bars(daily)
    if len(weekly) < 40:
        return None, "Insufficient completed weekly history"
    w_rsi = rsi(weekly["Close"], 5)
    w_rsi_sma = w_rsi.rolling(14, min_periods=14).mean()
    w_macd, w_signal = macd(weekly["Close"])
    w1 = pd.notna(w_rsi.iloc[-1]) and pd.notna(w_rsi_sma.iloc[-1]) and w_rsi.iloc[-1] > w_rsi_sma.iloc[-1]
    w2 = pd.notna(w_macd.iloc[-1]) and pd.notna(w_signal.iloc[-1]) and w_macd.iloc[-1] > w_signal.iloc[-1]

    # Monthly indicators: use completed monthly candles only.
    monthly = completed_monthly_bars(daily)
    if len(monthly) < 30:
        return None, "Insufficient completed monthly history"
    m_rsi = rsi(monthly["Close"], 5)
    m_rsi_sma = m_rsi.rolling(14, min_periods=14).mean()
    m_adx = adx(monthly["High"], monthly["Low"], monthly["Close"], 14)
    if len(monthly) < 2:
        return None, "Insufficient completed monthly candles"
    m1 = pd.notna(m_rsi.iloc[-1]) and pd.notna(m_rsi_sma.iloc[-1]) and m_rsi.iloc[-1] > m_rsi_sma.iloc[-1]
    last_adx = m_adx.iloc[-1]
    second_last_adx = m_adx.iloc[-2]
    adx_change = last_adx - second_last_adx if pd.notna(last_adx) and pd.notna(second_last_adx) else np.nan
    m2 = pd.notna(last_adx) and pd.notna(second_last_adx) and last_adx > second_last_adx
    m3 = pd.notna(adx_change) and 0 < adx_change < 1

    checks = [m1, m2, m3, w1, w2, d1, d2]
    qualified = all(checks)
    row = {
        "Symbol": symbol,
        "Company": company,
        "Yahoo Ticker": ticker,
        "Scan Date": daily.index[-1].date().isoformat(),
        "Close": daily_row["Close"],
        "Monthly RSI(5)": latest(m_rsi),
        "Monthly RSI SMA(14)": latest(m_rsi_sma),
        "Last Closed Monthly Candle": monthly.index[-1].date().isoformat(),
        "Monthly ADX(14) Last Closed": float(last_adx) if pd.notna(last_adx) else np.nan,
        "Monthly ADX(14) 2nd-Last Closed": float(second_last_adx) if pd.notna(second_last_adx) else np.nan,
        "Monthly ADX Change": float(adx_change) if pd.notna(adx_change) else np.nan,
        "Weekly RSI(5)": latest(w_rsi),
        "Weekly RSI SMA(14)": latest(w_rsi_sma),
        "Last Closed Weekly Candle": weekly.index[-1].date().isoformat(),
        "Weekly MACD": latest(w_macd),
        "Weekly MACD Signal": latest(w_signal),
        "Daily Low": daily_row["Daily Low"],
        "Daily EMA20": daily_row["EMA20"],
        "Daily MACD": daily_row["Daily MACD"],
        "Daily MACD Signal": daily_row["Daily MACD Signal"],
        "M1 Monthly RSI Pass": bool(m1),
        "M2 Monthly ADX Rising Pass": bool(m2),
        "M3 Monthly ADX Change 0-1 Pass": bool(m3),
        "W1 Weekly RSI Pass": bool(w1),
        "W2 Weekly MACD Pass": bool(w2),
        "D1 Daily EMA20 Pullback Pass": bool(d1),
        "D2 Daily MACD Pass": bool(d2),
        "QUALIFIES": bool(qualified),
    }
    return row, None


def style_excel(path: str):
    wb = load_workbook(path)
    ws = wb["Qualifying Stocks"]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 34
    for col_idx, col_cells in enumerate(ws.columns, 1):
        values = [str(c.value) if c.value is not None else "" for c in list(col_cells)[:200]]
        width = min(max(max((len(v) for v in values), default=10) + 2, 12), 34)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    # Add a simple instructions sheet.
    info = wb.create_sheet("Rules & Notes")
    notes = [
        ["Trial 3 — NIFTY 500 qualifying-stock list"],
        ["Monthly M1", "RSI(5) on last closed monthly candle > SMA(14) of monthly RSI(5) on that candle."],
        ["Monthly M2", "ADX(14) on last closed monthly candle > ADX(14) on second-last closed monthly candle."],
        ["Monthly M3", "0 < (last-closed monthly ADX(14) - second-last-closed monthly ADX(14)) < 1."],
        ["Weekly W1", "RSI(5) on last closed weekly candle > SMA(14) of weekly RSI(5) on that candle."],
        ["Weekly W2", "MACD(12,26,9) line > signal on the last closed weekly candle."],
        ["Daily D1", "Daily low <= EMA(20) and daily close > EMA(20), on the same daily candle."],
        ["Daily D2", "Daily MACD(12,26,9) line > signal on the daily candle being evaluated."],
        ["Important", "This is a stock screener, not a trade backtest. It does not calculate entries, stop-losses, targets, or exits."],
        ["Closed-candle policy", "The script conservatively excludes the current calendar month from monthly indicators. Weekly candles are included only when Friday's daily bar is present."],
        ["Data", "Yahoo Finance daily OHLCV and the NIFTY Indices constituent CSV. Data/provider delays and adjustments may affect results."],
        ["Universe note", "This scan uses the current NIFTY 500 constituent list, not historical point-in-time membership. Historical point-in-time membership is needed for unbiased historical backtests."],
    ]
    for row in notes:
        info.append(row)
    info.column_dimensions["A"].width = 25
    info.column_dimensions["B"].width = 115
    info.freeze_panes = "A2"
    for cell in info[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=14)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in info.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def main():
    now = datetime.now(IST)
    scan_date = now.date().isoformat()
    print(f"Trial 3 scan starting (IST): {now.isoformat()}")
    constituents = get_constituents()
    print(f"Loaded {len(constituents)} current NIFTY 500 constituents.")

    tickers = constituents["Ticker"].tolist()
    # Download in batches to reduce request failures and memory spikes.
    batch_size = 50
    all_rows = []
    failures = []
    for start in range(0, len(tickers), batch_size):
        batch = tickers[start:start + batch_size]
        print(f"Downloading batch {start // batch_size + 1}/{(len(tickers) + batch_size - 1) // batch_size}")
        try:
            data = yf.download(
                tickers=batch,
                period=DOWNLOAD_PERIOD,
                interval="1d",
                auto_adjust=False,
                group_by="ticker",
                progress=False,
                threads=True,
                timeout=30,
            )
        except Exception as e:
            failures.extend({"Ticker": t, "Reason": f"Batch download error: {e}"} for t in batch)
            continue

        for _, rec in constituents.iloc[start:start + batch_size].iterrows():
            ticker = rec["Ticker"]
            try:
                if isinstance(data.columns, pd.MultiIndex):
                    if ticker not in data.columns.get_level_values(0):
                        failures.append({"Ticker": ticker, "Reason": "No ticker data in batch response"})
                        continue
                    raw = data[ticker].copy()
                else:
                    raw = data.copy() if len(batch) == 1 else pd.DataFrame()
                row, error = evaluate_stock(rec["Symbol"], rec["Company"], ticker, raw, scan_date)
                if error:
                    failures.append({"Ticker": ticker, "Reason": error})
                elif row:
                    all_rows.append(row)
            except Exception as e:
                failures.append({"Ticker": ticker, "Reason": f"Evaluation error: {e}"})
        time.sleep(1)

    all_df = pd.DataFrame(all_rows)
    if all_df.empty:
        all_df = pd.DataFrame(columns=["Symbol", "Company", "Yahoo Ticker", "Scan Date", "QUALIFIES"])
        qualifying = all_df.copy()
    else:
        qualifying = all_df[all_df["QUALIFIES"] == True].copy()
        qualifying = qualifying.sort_values(["Symbol"]).reset_index(drop=True)

    # Excel: primary sheet contains only qualifying stocks, as requested.
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        qualifying.to_excel(writer, index=False, sheet_name="Qualifying Stocks")
        pd.DataFrame(failures, columns=["Ticker", "Reason"]).to_excel(writer, index=False, sheet_name="Download Issues")
    style_excel(OUTPUT_XLSX)
    qualifying.to_csv(OUTPUT_CSV, index=False)

    print(f"Evaluated: {len(all_df)}; qualifying: {len(qualifying)}; issues: {len(failures)}")
    print(f"Saved {OUTPUT_XLSX} and {OUTPUT_CSV}")

    # Optional Gmail delivery; existing GitHub Actions secrets can be used.
    username = os.getenv("GMAIL_USERNAME")
    app_password = os.getenv("GMAIL_APP_PASSWORD")
    recipient = os.getenv("GMAIL_TO")
    if username and app_password and recipient:
        import smtplib
        from email.message import EmailMessage
        msg = EmailMessage()
        msg["Subject"] = f"Trial 3 NIFTY 500 qualifying stocks — {scan_date} ({len(qualifying)})"
        msg["From"] = username
        msg["To"] = recipient
        if qualifying.empty:
            msg.set_content(f"No NIFTY 500 stocks qualified under Trial 3 on {scan_date}. The Excel file is attached.")
        else:
            preview = qualifying[["Symbol", "Company", "Close"]].head(40).to_string(index=False)
            msg.set_content(f"Trial 3 scan date: {scan_date}\nQualifying stocks: {len(qualifying)}\n\n{preview}\n\nFull list is attached in Excel.")
        with open(OUTPUT_XLSX, "rb") as f:
            msg.add_attachment(
                f.read(),
                maintype="application",
                subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                filename=OUTPUT_XLSX,
            )
        with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=30) as server:
            server.login(username, app_password)
            server.send_message(msg)
        print(f"Emailed report to {recipient}")
    else:
        print("Gmail secrets not all present; skipping email.")


if __name__ == "__main__":
    main()
